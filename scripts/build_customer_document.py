"""Build the customer-facing architecture, dataset, and Sabre MCP Word document."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Costco_Travel_Agentic_Demo_Technical_Guide.docx"
REPO = "https://github.com/ramamurthy-540835/costco-travel-demo"


def table(document, headers, rows):
    value = document.add_table(rows=1, cols=len(headers))
    value.style = "Light Shading Accent 1"
    for index, header in enumerate(headers):
        value.rows[0].cells[index].text = header
    for row in rows:
        cells = value.add_row().cells
        for index, item in enumerate(row):
            cells[index].text = str(item)
    return value


def bullets(document, values, level=0):
    for value in values:
        document.add_paragraph(value, style="List Bullet" if level == 0 else "List Bullet 2")


def steps(document, values):
    for value in values:
        document.add_paragraph(value, style="List Number")


def link_paragraph(document, label, url):
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(url)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.color.rgb = RGBColor(12, 41, 72)
    title = document.add_heading("Costco Travel Agentic Car Reservation Demo", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Dataset, Code Repository, Technology Stack, Delivery Flow, and Sabre MCP Integration")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Prepared {date.today().isoformat()} · Customer proof of concept · All member, price, and reservation records are synthetic").alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("1. Executive summary", level=1)
    document.add_paragraph("This proof of concept is a same-origin, single-service conversational rental-car experience. A browser UI and FastAPI backend run together on Cloud Run. Gemini interprets conversation, while deterministic server code—not the model—enforces date, cancellation, change, and human-review policies. Firestore stores demo reservations; GCS serves approved car imagery by signed URL; Secret Manager protects the SerpAPI credential; and a server-only SerpAPI adapter provides price discovery with a quota-safe fallback.")
    document.add_paragraph("The Sabre endpoint supplied for review is Sabre Developer Hub's documentation MCP server. It helps an AI agent discover and read Sabre API, SDK, and product-collection documentation. It is not, by itself, a live Sabre GDS shopping or booking endpoint. The repository now integrates it for documentation grounding and exposes a same-origin documentation-search API; transactional Sabre connectivity remains a separately provisioned phase.")

    document.add_heading("2. Code repository and deliverables", level=1)
    link_paragraph(document, "GitHub repository", REPO)
    table(document, ["Deliverable", "Repository location", "Purpose"], [
        ("Single-file frontend v3", "static/costco-travel-agent-v3.html", "Deployable UI with inline CSS/JS"),
        ("Backend", "app/", "FastAPI APIs, Gemini adapter, policies, reservations, search, images, Sabre MCP"),
        ("Synthetic bulk data", "docs/mock-data/", "500 rentals, 100 members, 1,000 bookings, providers, relative seed manifest"),
        ("Customer Word guide", "docs/Costco_Travel_Agentic_Demo_Technical_Guide.docx", "This document"),
        ("Tests", "tests/", "Policy, API, state-machine, SerpAPI, and Sabre MCP regressions"),
        ("Deployment", "Dockerfile and deploy.sh", "Non-root Cloud Run container and GCP provisioning"),
    ])

    document.add_heading("3. Dataset used by the demo", level=1)
    document.add_paragraph("There are two dataset layers. The operational layer is intentionally small and deterministic so the customer demo is repeatable. The bulk workshop layer reproduces the originally requested dataset sizes for analytics, demos, and future loading; it is not automatically inserted into Firestore.")
    table(document, ["Dataset", "Count", "Used in request path?", "Source / fields"], [
        ("Operational vehicle inventory", "6", "Yes", "app/recommend.py: car, class, seats, member rate, retail rate, popularity, emoji"),
        ("Relative Firestore reservations", "2", "Yes", "app/seed.py: 21 days out and 30 hours out; dates refresh when stale"),
        ("Frontend emergency fallback cars", "4", "Yes, only when /api/cars fails", "static/app.js: member/retail rates, vendor, features, savings"),
        ("Rental partners", "5", "Normalization and workshop", "Alamo, Avis, Budget, Enterprise, National"),
        ("Illustrative analytics aggregates", "5 providers + 5 locations", "Yes, insights screen", "Synthetic bookings and KPI values in app/main.py"),
        ("Bulk rental inventory", "500", "No; loadable fixture", "docs/mock-data/rental_inventory_500.json"),
        ("Bulk members", "100", "No; loadable fixture", "docs/mock-data/members_100.json; synthetic example.test emails; no credentials"),
        ("Bulk bookings", "1,000", "No; loadable fixture", "docs/mock-data/bookings_1000.json; future dates generated at build time"),
    ])
    document.add_heading("3.1 Operational vehicle inventory", level=2)
    table(document, ["Car", "Class", "Seats", "Member/day", "Retail/day"], [
        ("Toyota Camry", "Intermediate", 5, "$47", "$55"),
        ("Chrysler Pacifica", "Minivan", 7, "$69", "$82"),
        ("Chevrolet Tahoe", "Full-Size SUV", 7, "$86", "$101"),
        ("Toyota RAV4", "Standard SUV", 5, "$63", "$74"),
        ("Nissan Versa", "Economy", 5, "$36", "$43"),
        ("Chevrolet Malibu", "Full-Size", 5, "$54", "$64"),
    ])
    document.add_heading("3.2 Data classification and limitations", level=2)
    bullets(document, [
        "Every record and metric is synthetic and must remain labeled as demo/mock data.",
        "The demo presumes Costco SSO authentication. It does not collect passwords, membership credentials, card numbers, or income data.",
        "SerpAPI contributes public price-search observations only. Member rates are a deterministic 12–18% demonstration discount and are not a contractual Costco price.",
        "Car imagery comes only from the project GCS bucket and attribution manifest. Missing images return null and the emoji fallback remains visible.",
        "The bulk datasets are reproducible with scripts/generate_mock_data.py and should not be represented as Firestore production records unless deliberately loaded.",
    ])

    document.add_heading("4. Complete technology stack", level=1)
    table(document, ["Layer", "Technology", "Responsibility"], [
        ("Frontend", "HTML5, CSS3, vanilla JavaScript", "Preserved UX, modal calendar, cards, chips, sidebar, state machine"),
        ("Frontend packaging", "Single HTML build", "scripts/build_single_file.py produces costco-travel-agent-v3.html"),
        ("Backend", "Python 3.12, FastAPI, Uvicorn", "Same-origin static serving and /api routes"),
        ("Agent model", "Gemini 3 Flash Preview through google-genai", "Conversation reasoning and JSON action suggestions via Vertex AI ADC"),
        ("Deterministic policy", "Python pure functions", "Penalty tiers, tomorrow rule, HITL triggers, recommendation ranking"),
        ("Reservation store", "Firestore Native mode", "CONFIRMED/HOLD/PENDING/CANCELLED workflow state"),
        ("Images", "Cloud Storage + V4 signed URLs", "Approved class-keyed assets; ADC self-impersonated signing"),
        ("Secrets", "Secret Manager", "SerpAPI secret version 2; never sent to the browser or logged"),
        ("External inventory", "SerpAPI through server-side HTTP", "Price discovery, 15-minute cache, fallback on any failure"),
        ("Sabre knowledge", "Sabre Developer Hub MCP 1.0", "Search and retrieve Sabre documentation for grounding"),
        ("Runtime", "Cloud Run", "One public same-origin service, non-root container, 0–3 instances"),
        ("Build", "Cloud Build + Artifact Registry via gcloud run deploy --source", "Container build and revision deployment"),
        ("Analytics", "BigQuery optional; synthetic API aggregates today", "Event logging when BQ_ENABLED=true and executive demo metrics"),
        ("Identity", "Application Default Credentials and IAM", "No API keys or service-account JSON in source/image"),
        ("Testing", "pytest + Node syntax/state harness", "Server policies, HTTP flows, upstream fallback, UI transition invariants"),
    ])

    document.add_heading("5. End-to-end architecture", level=1)
    document.add_paragraph("Browser → Cloud Run (static v3 + FastAPI) → deterministic policy/orchestration → Firestore / GCS / Secret Manager / Vertex AI. The car-search adapter calls SerpAPI only after server date validation. The Sabre documentation adapter calls the public Sabre Developer Hub MCP only for documentation questions. Optional event records flow to BigQuery without affecting a transaction response.")
    document.add_heading("5.1 Search and new booking", level=2)
    steps(document, [
        "Member asks for a car. Client intent starts a new flow but never accepts typed dates as authoritative.",
        "Calendar opens once. Pickup minimum is tomorrow; defaults are today +7 and +11 days; return stays after pickup.",
        "Browser calls same-origin GET /api/cars with location and selected dates.",
        "Backend validates dates before Secret Manager or SerpAPI access, then checks the 15-minute cache.",
        "Results are normalized to member/retail price and savings. Any upstream failure returns sample inventory with source=fallback.",
        "Member selects a car and explicitly confirms the summary. POST /api/reservations creates a CONFIRMED Firestore record.",
    ])
    document.add_heading("5.2 Protected change", level=2)
    steps(document, [
        "POST change/start evaluates HITL. If safe, original becomes HOLD and a PENDING replacement shell is created atomically.",
        "The date picker appears once. change/update validates dates and cost delta while the original remains protected.",
        "Member selects a replacement and reviews old HOLD versus new PENDING, price delta, percentage, and savings.",
        "On explicit confirmation, replacement becomes CONFIRMED first. One microsecond later the original becomes CANCELLED.",
        "Abandonment deletes PENDING and restores HOLD to CONFIRMED. A cancelled reservation cannot be changed again.",
    ])
    document.add_heading("5.3 Two-step cancellation", level=2)
    steps(document, [
        "cancel/preview computes the exact tier, fee, refund, and free-cancel deadline without changing state.",
        "The UI renders one confirmation card with Confirm and Keep actions.",
        "Only the explicit confirm endpoint can mark a safe reservation CANCELLED. HITL returns 409 with no state change.",
        "Penalty schedule: at least 48 hours is free; 24–48 hours is one daily rate; under 24 hours is the no-show fee.",
    ])

    document.add_heading("6. Server-side guardrails", level=1)
    bullets(document, [
        "Pickup must be tomorrow or later in the pickup timezone; return must follow pickup. Create, change, and car search return HTTP 422 when invalid.",
        "HITL: change within 48 hours, cancellation within 24 hours, same-day action, penalty exposure over $100, or total cost change over 30%.",
        "Model responses never directly mutate reservations. All mutations are endpoint-controlled and policy-checked.",
        "Cancellation always has preview and confirm. Change always uses HOLD → PENDING → explicit confirm.",
        "No credentials, payment card numbers, or income data are collected; no return/income guarantees are made.",
    ])

    document.add_heading("7. Sabre MCP Server analysis and integration", level=1)
    document.add_paragraph("Direct protocol inspection of https://developer.mcp.sabre.com/mcp returned server name dev-studio-devhub-mcp, version 1.0.0, MCP protocol 2025-03-26, and tools/resources/prompts capabilities. Its server instructions describe a Sabre Developer Hub documentation assistant.")
    table(document, ["Official tool", "What it does"], [
        ("look_for_named_sabre_documentation", "Finds APIs, SDKs, and product collections by artifact/service name"),
        ("look_for_artifact_content", "Searches documentation content and returns matching artifacts/snippets"),
        ("get_latest_updated_documentation", "Lists recently updated Sabre documentation"),
        ("get_artifact_details", "Returns artifact metadata and documentation-page paths"),
        ("get_documentation_page", "Retrieves a selected documentation page as Markdown or HTML"),
    ])
    document.add_paragraph("A live documentation search for 'rental car availability API' returned Sabre Car Availability REST 2.4.1, Cars Product Collection v1, Get Vehicle Availability REST v1/v2, the Car Reservation guide, and Book Car Reservation SOAP 2.2.0. These are candidate artifacts for the transactional discovery phase; customer entitlement and the current Sabre contract determine which one can be used.")
    document.add_heading("7.1 What is integrated now", level=2)
    bullets(document, [
        "app/sabre_mcp.py implements a bounded, server-side JSON-RPC tools/call client.",
        "GET /api/sabre/docs/search?q=... calls look_for_artifact_content through the same Cloud Run origin.",
        "Sabre-related chat questions retrieve documentation context and add it to Gemini's live system prompt.",
        "The prompt explicitly forbids claiming that documentation MCP performed a live shop, book, ticket, or service action.",
        "The endpoint requires no Sabre credential today because Sabre publishes this documentation MCP endpoint publicly.",
    ])
    document.add_heading("7.2 Transactional Sabre next phase", level=2)
    steps(document, [
        "Confirm the exact Sabre Cars shopping, pricing, booking, and servicing products licensed for the customer/PCC. Documentation MCP can help identify their artifacts.",
        "Obtain Sabre sandbox/production credentials contractually and store them only in Google Secret Manager. Never expose them in HTML, Gemini messages, logs, or the container.",
        "Add a Cloud Run Sabre tool adapter with OAuth/token caching, request schemas, response normalization, quotas, timeouts, idempotency keys, and redacted audit logs.",
        "Expose read-only shop/price tools first. Continue to use the existing deterministic recommendation and member-savings presentation.",
        "Place create/change/cancel Sabre calls behind the existing explicit confirmation and HITL state machine. The LLM may propose a tool, but server policy authorizes execution.",
        "Map a Sabre confirmation/record locator to the Costco demo reservation only after Sabre confirms success; use compensation/reconciliation for partial failures.",
        "Complete Sabre certification, PCI/privacy review, load tests, alerting, and a controlled production rollout.",
    ])

    document.add_heading("8. API surface", level=1)
    table(document, ["Method", "Path", "Purpose"], [
        ("GET", "/api/health", "Service, ADC, and model status"),
        ("POST", "/api/agent/chat", "Gemini JSON conversation response with deterministic fallback"),
        ("GET", "/api/cars", "Validated SerpAPI/fallback member inventory"),
        ("GET/POST", "/api/reservations...", "Create/list/get plus protected change and cancellation endpoints"),
        ("GET", "/api/policies/cancellation", "Machine-readable schedule and deadlines"),
        ("GET", "/api/sabre/docs/search", "Sabre Developer Hub documentation MCP search"),
        ("GET", "/api/analytics", "Synthetic executive-review metrics"),
    ])

    document.add_heading("9. Build, test, and deploy flow", level=1)
    steps(document, [
        "Authenticate locally: gcloud auth application-default login and gcloud auth login.",
        "Generate mock fixtures: python scripts/generate_mock_data.py.",
        "Build the single frontend: python scripts/build_single_file.py.",
        "Run pytest -q and compile/syntax checks. Run the credential-pattern scan before commit.",
        "Run locally with RESERVATION_BACKEND=memory for an offline demo, or ADC + Firestore for the cloud-equivalent path.",
        "Execute deploy.sh. It enables APIs, provisions IAM/GCS, grants Secret Manager access, and deploys Cloud Run using the non-root Dockerfile.",
        "Validate health, cars, chat JSON, past-date 422, protected change ordering, HITL untouched state, penalty tiers, image URL/null, and the browser flow.",
        "Commit and push only code, generated synthetic fixtures, documentation, and tests—never local environment files or secret values.",
    ])

    document.add_heading("10. Sources reviewed", level=1)
    link_paragraph(document, "Sabre MCP product link supplied by customer", "https://developer.sabre.com/product-collection/mcp-server/1.0/index.html")
    link_paragraph(document, "Sabre Developer Hub documentation MCP endpoint", "https://developer.mcp.sabre.com/mcp")
    link_paragraph(document, "Sabre agentic API/MCP announcement", "https://investors.sabre.com/news-releases/news-release-details/sabre-seizes-first-mover-position-comprehensive-agentic-apis")
    link_paragraph(document, "Google Vertex AI Gemini quickstart/model examples", "https://docs.cloud.google.com/vertex-ai/generative-ai/docs/start/quickstart")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(f"built {OUT}")


if __name__ == "__main__":
    main()
